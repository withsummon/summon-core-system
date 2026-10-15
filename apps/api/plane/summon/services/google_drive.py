# Copyright (c) 2023-present Plane Software, Inc. and contributors
# SPDX-License-Identifier: AGPL-3.0-only
# See the LICENSE file for details.

"""Google Drive folders synced into project pages through a member's Google account.

A member connects their Google account once (OAuth, read-only Drive access); the refresh token lives in the
credential vault. Connecting a folder copies its Docs, Sheets and Slides into locked pages; connecting a single
Doc, Sheet or Slides file copies just that one. After that Drive's
change feed keeps them current: a push channel tells us something changed and `apply_drive_changes` reads what,
while a periodic run of the same function catches anything a notification missed.
"""

import os
import re
import secrets
import time
from datetime import timedelta
from html import escape
from io import BytesIO
from urllib.parse import urlencode
from uuid import uuid4

import jwt
import openpyxl
import requests
from django.conf import settings
from django.db import transaction
from django.utils import timezone

from plane.db.models import Page, ProjectPage
from plane.license.utils.instance_value import get_configuration_value
from plane.summon.models import Credential, DriveFile, DriveFolder
from plane.summon.services.credential import decrypt_secret, encrypt_secret
from plane.summon.services.page_document import write_page_document

GOOGLE_DRIVE_PROVIDER = "google_drive"
AUTH_URL = "https://accounts.google.com/o/oauth2/v2/auth"
TOKEN_URL = "https://oauth2.googleapis.com/token"
DRIVE_API_URL = "https://www.googleapis.com/drive/v3"
# Docs are read through the Docs API rather than exported: Drive refuses to export a Doc whose HTML, with every
# image inlined, passes 10 MB, and its export only holds the first tab. drive.readonly covers this read too.
DOCS_API_URL = "https://docs.googleapis.com/v1"
# Only what a page keeps (text, emphasis, links, headings, lists, tables); the full document carries every style.
DOC_CONTENT_FIELDS = (
    "content(paragraph(elements/textRun(content,textStyle(bold,italic,link/url)),"
    "paragraphStyle/namedStyleType,bullet(listId,nestingLevel)),table/tableRows/tableCells/content)"
)
DOC_TAB_FIELDS = f"tabProperties/title,documentTab(lists,body/{DOC_CONTENT_FIELDS})"
DOC_FIELDS = f"tabs({DOC_TAB_FIELDS},childTabs({DOC_TAB_FIELDS},childTabs({DOC_TAB_FIELDS})))"
DRIVE_SCOPE = "https://www.googleapis.com/auth/drive.readonly"
OAUTH_SCOPE = f"openid email {DRIVE_SCOPE}"
DRIVE_TIMEOUT_SECONDS = 30
FOLDER_MIME = "application/vnd.google-apps.folder"
FILE_FIELDS = "id,name,mimeType,modifiedTime,parents,trashed,webViewLink"
# Drive caps a changes channel at a week; renew a day early so a slow beat never leaves a gap.
CHANNEL_LIFETIME = timedelta(days=7)
CHANNEL_RENEW_WITHIN = timedelta(days=1)
# A synced sheet is reference material, not a database: past these it is cut, and the page says so.
SHEET_MAX_ROWS = 500
SHEET_MAX_COLUMNS = 30

# Folder links (/folders/<id>), file links (/document/d/<id>, /spreadsheets/d/<id>, /file/d/<id>, …) and
# the older ?id=<id> form all carry the Drive ID; Drive itself says which kind it is.
_LINKED_ID = re.compile(r"(?:/folders/|/d/|[?&]id=)([\w-]{10,})")
_BARE_ID = re.compile(r"^[\w-]{10,}$")


class GoogleDriveError(Exception):
    def __init__(self, code):
        super().__init__(code)
        self.code = code


def parse_drive_id(value):
    """The Drive ID in a folder or file link as Drive and Docs show it, or a bare ID."""
    value = value.strip()
    match = _LINKED_ID.search(value)
    if match:
        return match.group(1)
    if _BARE_ID.match(value):
        return value
    raise GoogleDriveError("drive_link_invalid")


def webhook_address():
    """Drive only notifies public HTTPS addresses; elsewhere (local dev) the periodic run keeps folders current."""
    base = settings.WEB_URL or ""
    return f"{base.rstrip('/')}/api/summon/google-drive/notifications/" if base.startswith("https://") else None


# Google account --------------------------------------------------------------------------------------------


def _oauth_client():
    # The OAuth client Plane's "Sign in with Google" uses; its consent screen must also allow Drive read access.
    client_id, client_secret = get_configuration_value(
        [
            {"key": "GOOGLE_CLIENT_ID", "default": os.environ.get("GOOGLE_CLIENT_ID")},
            {"key": "GOOGLE_CLIENT_SECRET", "default": os.environ.get("GOOGLE_CLIENT_SECRET")},
        ]
    )
    if not (client_id and client_secret):
        raise GoogleDriveError("google_not_configured")
    return client_id, client_secret


def authorization_url(redirect_uri, state):
    client_id, _ = _oauth_client()
    params = {
        "client_id": client_id,
        "redirect_uri": redirect_uri,
        "response_type": "code",
        "scope": OAUTH_SCOPE,
        "state": state,
        # Google only hands out a refresh token for offline access, and reliably only on a fresh consent.
        "access_type": "offline",
        "prompt": "consent",
    }
    return f"{AUTH_URL}?{urlencode(params)}"


def _token_request(data):
    client_id, client_secret = _oauth_client()
    try:
        response = requests.post(
            TOKEN_URL,
            data={**data, "client_id": client_id, "client_secret": client_secret},
            timeout=DRIVE_TIMEOUT_SECONDS,
        )
        body = response.json()
    except (requests.RequestException, ValueError):
        raise GoogleDriveError("drive_unavailable") from None
    if response.status_code == 200:
        return body
    if body.get("error") in ("invalid_client", "unauthorized_client"):
        raise GoogleDriveError("google_not_configured")
    if body.get("error") == "invalid_grant":
        # The person revoked access, changed their password, or an admin removed the app.
        raise GoogleDriveError("drive_auth_revoked")
    raise GoogleDriveError("drive_unavailable")


def exchange_code(code, redirect_uri):
    """The Google account behind an authorization code: its address, and a refresh token for reading Drive."""
    body = _token_request({"grant_type": "authorization_code", "code": code, "redirect_uri": redirect_uri})
    # The consent screen lets people untick Drive and still finish.
    if DRIVE_SCOPE not in body.get("scope", "").split() or not body.get("refresh_token"):
        raise GoogleDriveError("drive_scope_missing")
    # The ID token came straight from Google's token endpoint over TLS, which OpenID Connect accepts without
    # checking its signature.
    email = jwt.decode(body["id_token"], options={"verify_signature": False})["email"]
    return email, body["refresh_token"]


def save_google_account(workspace, member, email, refresh_token):
    """Keep a member's Google account in the vault, reactivating it when they reconnect after a revocation."""
    credential = Credential.objects.filter(
        workspace=workspace, owner=member, provider=GOOGLE_DRIVE_PROVIDER, account_identifier=email
    ).first() or Credential(
        workspace=workspace,
        owner=member,
        provider=GOOGLE_DRIVE_PROVIDER,
        account_identifier=email,
        name=f"Google Drive ({email})",
    )
    credential.secret_ciphertext = encrypt_secret(refresh_token)
    credential.status = Credential.Status.ACTIVE
    credential.revoked_at = None
    credential.save()
    return credential


class DriveClient:
    def __init__(self, refresh_token):
        self._refresh_token = refresh_token
        self._token = None
        self._token_expires = 0
        self.session = requests.Session()

    def _access_token(self):
        if self._token and self._token_expires > time.time() + 60:
            return self._token
        body = _token_request({"grant_type": "refresh_token", "refresh_token": self._refresh_token})
        self._token = body["access_token"]
        self._token_expires = time.time() + int(body.get("expires_in", 3600))
        return self._token

    def _request(self, method, path, api=DRIVE_API_URL, **kwargs):
        try:
            response = self.session.request(
                method,
                f"{api}{path}",
                headers={"Authorization": f"Bearer {self._access_token()}"},
                timeout=DRIVE_TIMEOUT_SECONDS,
                **kwargs,
            )
        except requests.RequestException:
            raise GoogleDriveError("drive_unavailable") from None
        if response.status_code == 404:
            # Drive answers 404 for anything the account cannot open, whether or not it exists.
            raise GoogleDriveError("drive_not_found")
        if response.status_code == 401:
            raise GoogleDriveError("drive_auth_revoked")
        if response.status_code == 403:
            raise GoogleDriveError("drive_access_denied")
        if response.status_code >= 300:
            raise GoogleDriveError("drive_unavailable")
        return response

    def item(self, item_id):
        return self._request(
            "GET",
            f"/files/{item_id}",
            params={"fields": f"{FILE_FIELDS},driveId", "supportsAllDrives": "true"},
        ).json()

    def children(self, folder_id):
        params = {
            "q": f"'{folder_id}' in parents and trashed = false",
            "fields": f"nextPageToken,files({FILE_FIELDS})",
            "pageSize": 1000,
            "corpora": "allDrives",
            "supportsAllDrives": "true",
            "includeItemsFromAllDrives": "true",
        }
        while True:
            body = self._request("GET", "/files", params=params).json()
            yield from body.get("files", [])
            if not body.get("nextPageToken"):
                return
            params["pageToken"] = body["nextPageToken"]

    def walk(self, folder_id):
        """Every folder under `folder_id` (itself included) and every file in them."""
        folder_ids, files, pending = [folder_id], [], [folder_id]
        while pending:
            for item in self.children(pending.pop()):
                if item["mimeType"] == FOLDER_MIME:
                    folder_ids.append(item["id"])
                    pending.append(item["id"])
                else:
                    files.append(item)
        return folder_ids, files

    def _drive_params(self, drive_id):
        params = {"supportsAllDrives": "true"}
        if drive_id:
            params["driveId"] = drive_id
        return params

    def start_page_token(self, drive_id):
        return self._request("GET", "/changes/startPageToken", params=self._drive_params(drive_id)).json()[
            "startPageToken"
        ]

    def changes(self, page_token, drive_id):
        """The changes since `page_token`, and the token to read from next time."""
        params = {
            **self._drive_params(drive_id),
            "pageToken": page_token,
            "pageSize": 1000,
            "includeItemsFromAllDrives": "true",
            "includeRemoved": "true",
            "fields": f"nextPageToken,newStartPageToken,changes(fileId,removed,file({FILE_FIELDS}))",
        }
        changes = []
        while True:
            body = self._request("GET", "/changes", params=params).json()
            changes.extend(body.get("changes", []))
            if body.get("newStartPageToken"):
                return changes, body["newStartPageToken"]
            params["pageToken"] = body["nextPageToken"]

    def export(self, file_id, mime_type):
        return self._request("GET", f"/files/{file_id}/export", params={"mimeType": mime_type}).content

    def document(self, file_id):
        return self._request(
            "GET",
            f"/documents/{file_id}",
            api=DOCS_API_URL,
            params={"includeTabsContent": "true", "fields": DOC_FIELDS},
        ).json()

    def watch(self, page_token, drive_id, channel_id, channel_token, address, expires_at):
        return self._request(
            "POST",
            "/changes/watch",
            params={**self._drive_params(drive_id), "pageToken": page_token, "includeItemsFromAllDrives": "true"},
            json={
                "id": channel_id,
                "type": "web_hook",
                "address": address,
                "token": channel_token,
                "expiration": int(expires_at.timestamp() * 1000),
            },
        ).json()

    def stop(self, channel_id, resource_id):
        self._request("POST", "/channels/stop", json={"id": channel_id, "resourceId": resource_id})


def drive_client(credential):
    if credential.status != Credential.Status.ACTIVE:
        raise GoogleDriveError("drive_auth_revoked")
    return DriveClient(decrypt_secret(credential.secret_ciphertext))


def record_drive_error(folder, error):
    folder.sync_error = error.code
    folder.save(update_fields=["sync_error", "updated_at"])
    if error.code == "drive_auth_revoked" and folder.credential.status == Credential.Status.ACTIVE:
        # Every folder on this account is broken now, so the account itself is marked; reconnecting restores it.
        folder.credential.status = Credential.Status.REVOKED
        folder.credential.revoked_at = timezone.now()
        folder.credential.save(update_fields=["status", "revoked_at", "updated_at"])


# Content -----------------------------------------------------------------------------------------------------


_HEADINGS = {"TITLE": "h1", "SUBTITLE": "h2", **{f"HEADING_{level}": f"h{level}" for level in range(1, 7)}}
_LINK_SCHEMES = ("https://", "http://", "mailto:")


def _doc_html(client, file_id):
    tabs = list(_doc_tabs(client.document(file_id)["tabs"]))
    # A Doc with several tabs puts each under its title.
    titled = len(tabs) > 1
    return "".join(
        (f"<h1>{escape(tab['tabProperties']['title'])}</h1>" if titled else "")
        + _content_html(tab["documentTab"]["body"]["content"], tab["documentTab"].get("lists", {}))
        for tab in tabs
    )


def _doc_tabs(tabs):
    for tab in tabs:
        yield tab
        yield from _doc_tabs(tab.get("childTabs", []))


def _content_html(content, lists):
    """A Doc's structural elements as page HTML. Bullets nest by their level; a table cell holds its own content."""
    parts, open_lists = [], []
    for element in content:
        paragraph = element.get("paragraph")
        bullet = (paragraph or {}).get("bullet")
        depth = bullet.get("nestingLevel", 0) + 1 if bullet else 0
        while len(open_lists) > depth:
            parts.append(f"</li></{open_lists.pop()}>")
        if bullet:
            if len(open_lists) == depth:
                parts.append("</li>")
            while len(open_lists) < depth:
                open_lists.append(_list_tag(lists, bullet["listId"], len(open_lists)))
                parts.append(f"<{open_lists[-1]}>")
            parts.append(f"<li><p>{_runs_html(paragraph)}</p>")
        elif paragraph:
            text = _runs_html(paragraph)
            tag = _HEADINGS.get(paragraph.get("paragraphStyle", {}).get("namedStyleType"), "p")
            # Docs spaces paragraphs with empty ones; the page's own spacing replaces them.
            if text.strip():
                parts.append(f"<{tag}>{text}</{tag}>")
        elif "table" in element:
            rows = "".join(
                "<tr>"
                + "".join(
                    f"<td>{_content_html(cell.get('content', []), lists) or '<p></p>'}</td>"
                    for cell in row.get("tableCells", [])
                )
                + "</tr>"
                for row in element["table"].get("tableRows", [])
            )
            parts.append(f"<table><tbody>{rows}</tbody></table>")
    parts.extend(f"</li></{tag}>" for tag in reversed(open_lists))
    return "".join(parts)


def _list_tag(lists, list_id, level):
    levels = lists.get(list_id, {}).get("listProperties", {}).get("nestingLevels", [])
    glyph = levels[level].get("glyphType", "") if level < len(levels) else ""
    # Numbered levels name a glyph type (DECIMAL, ALPHA, …); bulleted ones carry a symbol instead.
    return "ol" if glyph and glyph != "GLYPH_TYPE_UNSPECIFIED" else "ul"


def _runs_html(paragraph):
    parts = []
    for element in paragraph.get("elements", []):
        run = element.get("textRun")
        if not run:
            continue
        # A run ends its paragraph with "\n"; "\x0b" is a line break inside one, and U+E907 stands in for an
        # image, chip or other object the Docs API does not return as text.
        text = escape(run["content"].rstrip("\n").replace("\ue907", "")).replace("\x0b", "<br>")
        if not text:
            continue
        style = run.get("textStyle", {})
        if style.get("bold"):
            text = f"<strong>{text}</strong>"
        if style.get("italic"):
            text = f"<em>{text}</em>"
        url = style.get("link", {}).get("url", "")
        if url.startswith(_LINK_SCHEMES):
            text = f'<a href="{escape(url)}">{text}</a>'
        parts.append(text)
    return "".join(parts)


def _sheet_html(client, file_id):
    workbook = openpyxl.load_workbook(
        BytesIO(client.export(file_id, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")),
        read_only=True,
        data_only=True,
    )
    parts = []
    for sheet in workbook.worksheets:
        rows = []
        for row in sheet.iter_rows(values_only=True, max_row=SHEET_MAX_ROWS):
            cells = ["" if value is None else str(value) for value in row[:SHEET_MAX_COLUMNS]]
            while cells and not cells[-1]:
                cells.pop()
            rows.append(cells)
        while rows and not rows[-1]:
            rows.pop()
        if not rows:
            continue
        width = max(len(row) for row in rows)
        table = "".join(
            "<tr>" + "".join(f"<td><p>{escape(cell)}</p></td>" for cell in row + [""] * (width - len(row))) + "</tr>"
            for row in rows
        )
        cut = (sheet.max_row or 0) > SHEET_MAX_ROWS or (sheet.max_column or 0) > SHEET_MAX_COLUMNS
        note = f"<p><em>First {SHEET_MAX_ROWS} rows and {SHEET_MAX_COLUMNS} columns.</em></p>" if cut else ""
        parts.append(f"<h2>{escape(sheet.title)}</h2>{note}<table><tbody>{table}</tbody></table>")
    workbook.close()
    return "".join(parts)


def _slides_html(client, file_id):
    text = client.export(file_id, "text/plain").decode("utf-8", "replace")
    return "".join(f"<p>{escape(line)}</p>" for line in text.splitlines() if line.strip())


CONVERTERS = {
    "application/vnd.google-apps.document": _doc_html,
    "application/vnd.google-apps.spreadsheet": _sheet_html,
    "application/vnd.google-apps.presentation": _slides_html,
}


# Sync --------------------------------------------------------------------------------------------------------


def _write_file(folder, client, item):
    """Copy one Drive file into its page, unless the page already holds this revision."""
    record = DriveFile.objects.filter(folder=folder, file_id=item["id"]).select_related("page").first()
    if record and record.modified_time == item["modifiedTime"]:
        return
    page = record.page if record else Page(workspace=folder.workspace, owned_by=folder.created_by, is_global=False)
    page.name = item["name"][:255]
    page.access = Page.PUBLIC_ACCESS
    # Drive owns the content; edits made here would be overwritten by the next change.
    page.is_locked = True
    write_page_document(
        page,
        # Plane Live rejects HTML without a tag, which an empty file would be.
        CONVERTERS[item["mimeType"]](client, item["id"]) or "<p></p>",
        {
            "kind": "google_drive",
            "file_id": item["id"],
            "url": item.get("webViewLink", ""),
            "modified_time": item["modifiedTime"],
        },
    )
    ProjectPage.objects.get_or_create(workspace=folder.workspace, project=folder.project, page=page)
    if record:
        record.modified_time = item["modifiedTime"]
        record.save(update_fields=["modified_time", "updated_at"])
    else:
        DriveFile.objects.create(
            workspace=folder.workspace,
            folder=folder,
            page=page,
            file_id=item["id"],
            modified_time=item["modifiedTime"],
        )


def _archive_files(folder, file_ids):
    """Files gone from the folder leave their pages archived, not deleted, in case they come back."""
    for record in DriveFile.objects.filter(folder=folder, file_id__in=file_ids).select_related("page"):
        record.page.archived_at = timezone.now().date()
        record.page.save(update_fields=["archived_at", "updated_at"])
        record.delete(soft=False)


def _finish(folder, **fields):
    for name, value in fields.items():
        setattr(folder, name, value)
    folder.last_synced_at = timezone.now()
    folder.sync_error = ""
    folder.save(update_fields=[*fields, "last_synced_at", "sync_error", "updated_at"])


def sync_drive_folder(folder, client=None):
    """Bring every page in line with the folder as it is now: the first import, and the repair for anything the
    change feed cannot express on its own, such as a whole folder moving in."""
    client = client or drive_client(folder.credential)
    # Read the feed position before listing, so a change made during the walk is read again rather than lost.
    page_token = folder.page_token or client.start_page_token(folder.drive_id)
    if folder.kind == DriveFolder.Kind.FILE:
        item = client.item(folder.folder_id)
        folder_ids, files, name = [], [] if item.get("trashed") else [item], item["name"]
    else:
        (folder_ids, files), name = client.walk(folder.folder_id), folder.name
    supported = [item for item in files if item["mimeType"] in CONVERTERS]
    for item in supported:
        _write_file(folder, client, item)
    present = {item["id"] for item in supported}
    _archive_files(folder, set(folder.files.values_list("file_id", flat=True)) - present)
    _finish(folder, folder_ids=folder_ids, page_token=page_token, name=name)


def apply_drive_changes(folder, client=None):
    """Apply what changed in Drive since the folder's feed position."""
    client = client or drive_client(folder.credential)
    changes, next_token = client.changes(folder.page_token, folder.drive_id)
    tree = set(folder.folder_ids)
    tracked = set(folder.files.values_list("file_id", flat=True))
    written, gone = [], set()
    for change in changes:
        item = change.get("file") or {}
        # A file source has no tree: the file itself is all that belongs to it.
        belongs = bool(tree & set(item.get("parents", []))) or change["fileId"] == folder.folder_id
        inside = not change.get("removed") and not item.get("trashed") and belongs
        if item.get("mimeType") == FOLDER_MIME and (inside or item.get("id") in tree):
            # A folder appearing, moving or vanishing reshapes the tree; walking it again is simpler than patching.
            folder.page_token = next_token
            return sync_drive_folder(folder, client)
        if inside and item.get("mimeType") in CONVERTERS:
            written.append(item)
        elif change["fileId"] in tracked:
            gone.add(change["fileId"])
    for item in written:
        _write_file(folder, client, item)
    _archive_files(folder, gone)
    # A synced file is listed under its own name, so a rename there renames the source too.
    renamed = {item["name"] for item in written if item["id"] == folder.folder_id}
    _finish(folder, page_token=next_token, **({"name": renamed.pop()} if renamed else {}))


def run_drive_sync(folder_id, full=False):
    """One sync at a time per folder: notifications arrive in bursts, and two runs would race to create pages."""
    with transaction.atomic():
        folder = (
            DriveFolder.objects.select_for_update(of=("self",))
            .select_related("credential", "workspace", "project", "created_by")
            .filter(id=folder_id)
            .first()
        )
        if not folder:
            return None
        try:
            if full or not folder.page_token:
                sync_drive_folder(folder)
            else:
                apply_drive_changes(folder)
        except GoogleDriveError as error:
            record_drive_error(folder, error)
        return folder


def watch_drive_folder(folder, client=None):
    """Open a push channel for the folder's feed, replacing the one it had. Without a public address, skip."""
    address = webhook_address()
    if not address:
        return
    client = client or drive_client(folder.credential)
    expires_at = timezone.now() + CHANNEL_LIFETIME
    channel_id, channel_token = uuid4().hex, secrets.token_urlsafe(32)
    channel = client.watch(folder.page_token, folder.drive_id, channel_id, channel_token, address, expires_at)
    previous = (folder.channel_id, folder.channel_resource_id)
    folder.channel_id = channel_id
    folder.channel_token = channel_token
    folder.channel_resource_id = channel["resourceId"]
    folder.channel_expires_at = expires_at
    folder.save(
        update_fields=["channel_id", "channel_token", "channel_resource_id", "channel_expires_at", "updated_at"]
    )
    if all(previous):
        stop_drive_channel(client, *previous)


def stop_drive_channel(client, channel_id, resource_id):
    try:
        client.stop(channel_id, resource_id)
    except GoogleDriveError:
        pass  # Already expired or stopped; an unknown channel's notifications are ignored anyway.


def disconnect_drive_folder(folder):
    """Stop listening and keep the pages, unlocked, as ordinary pages people can now edit."""
    if folder.channel_id:
        try:
            stop_drive_channel(drive_client(folder.credential), folder.channel_id, folder.channel_resource_id)
        except GoogleDriveError:
            pass  # A revoked account cannot stop its channel; the channel expires and its notifications are ignored.
    Page.objects.filter(summon_drive_files__folder=folder).update(is_locked=False)
    # Sync bookkeeping, not anyone's content: gone for good, files with it.
    folder.delete(soft=False)
